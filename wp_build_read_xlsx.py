import openpyxl, sys

path = r"C:/Users/liush/Desktop/Ai/wordpress/HOLANG网站资料总表-260720.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("SHEETS:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== SHEET: {name}  (rows={ws.max_row}, cols={ws.max_column}) ===")
    rows = list(ws.iter_rows(values_only=True))
    # print header-ish first 3 rows
    for i, r in enumerate(rows[:3]):
        # trim None
        cells = [str(c)[:40] if c is not None else "" for c in r]
        print(f"  r{i+1}:", " | ".join(cells[:12]))
    if len(rows) > 3:
        print(f"  ... total {len(rows)} rows")
wb.close()
